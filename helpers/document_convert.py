#!/usr/bin/env python3
"""Private, bounded document worker. JSON diagnostics never contain document data."""

from __future__ import annotations

import argparse
import base64
from contextlib import contextmanager
import html
import importlib.util
import io
import json
import math
import mimetypes
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET


MAX_SOURCE = 64 * 1024 * 1024
MAX_OUTPUT = 256 * 1024 * 1024
MAX_TEXT = 16 * 1024 * 1024
MAX_PAGES = 200
DOCUMENT_EXTENSIONS = {"txt", "md", "markdown", "html", "htm", "docx", "odt", "rtf"}
DOCUMENT_FORMATS = ["pdf", "docx", "odt", "rtf", "txt", "md", "html", "epub"]
FILTERS = {
    "pdf": "pdf:writer_pdf_Export",
    "docx": "docx:Office Open XML Text",
    "odt": "odt:writer8",
    "rtf": "rtf:Rich Text Format",
    "txt": "txt:Text (encoded):UTF8",
    "md": "md:Markdown",
    "html": "html:HTML (StarWriter)",
    "epub": "epub:EPUB",
}


class ConversionError(Exception):
    """Messages are fixed product strings, never untrusted library diagnostics."""


_OCR_MODULE = None
_PLATFORM_MODULE = None
_RUNTIME = None
_RUNTIME_ROOT = ""


def platform_module():
    global _PLATFORM_MODULE
    if _PLATFORM_MODULE is None:
        path = Path(__file__).resolve().with_name("platform_runtime.py")
        specification = importlib.util.spec_from_file_location("maracuja_private_platform", path)
        if specification is None or specification.loader is None:
            raise ConversionError("O componente local de conversão não está disponível.")
        module = importlib.util.module_from_spec(specification)
        specification.loader.exec_module(module)
        _PLATFORM_MODULE = module
    return _PLATFORM_MODULE


def runtime():
    global _RUNTIME
    if _RUNTIME is None:
        _RUNTIME = platform_module().Runtime(_RUNTIME_ROOT)
    return _RUNTIME


def required_tool(name: str) -> str:
    tool = runtime().tool(name)
    if tool is None:
        raise ConversionError("Um componente local necessário para converter PDF não está disponível.")
    return str(tool)


def ocr_module():
    global _OCR_MODULE
    if _OCR_MODULE is None:
        # Python -I excludes the script's directory from sys.path. Load only this
        # installed sibling, never modules from the source document's directory.
        path = Path(__file__).resolve().with_name("ocr_engine.py")
        specification = importlib.util.spec_from_file_location("maracuja_private_ocr", path)
        if specification is None or specification.loader is None:
            raise ConversionError("O componente local de OCR não está disponível.")
        module = importlib.util.module_from_spec(specification)
        specification.loader.exec_module(module)
        _OCR_MODULE = module
    return _OCR_MODULE


def emit_progress(done: int, total: int, phase: str) -> None:
    print(json.dumps({"type": "progress", "done": done, "total": total, "phase": phase}), flush=True)


def clean_text(text: str) -> str:
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", "", text)


def has_module(name: str) -> bool:
    try:
        return importlib.util.find_spec(name) is not None
    except (ImportError, ValueError):
        return False


def office_root(executable: str) -> Path | None:
    return platform_module().office_root(runtime().tool("soffice", executable))


def office_isolation_tool() -> Path | None:
    executable = runtime().tool("bwrap")
    return executable if executable is not None and os.access(executable, os.X_OK) else None


def capabilities(soffice: str, tessdata: str = "") -> dict:
    root = office_root(soffice)
    available = root is not None and (runtime().windows or office_isolation_tool() is not None)
    document_formats = list(DOCUMENT_FORMATS)
    if not platform_module().markdown_filter(root):
        document_formats.remove("md")
    pdf = []
    if has_module("pypdf"):
        if has_module("docx"):
            pdf.append("docx")
        pdf.extend(["txt", "md", "html", "epub"])
        if has_module("pptx") and has_module("PIL") and runtime().tool("pdftoppm") is not None:
            pdf.append("pptx")
        if has_module("pdfplumber") and has_module("openpyxl"):
            pdf.append("xlsx")
    try:
        ocr = ocr_module().capabilities(tessdata, _RUNTIME_ROOT)
    except Exception:
        ocr = {"ocr": False, "langs": []}
    image_formats = ["txt"] if ocr["ocr"] else []
    if image_formats and has_module("docx"):
        image_formats.append("docx")
    return {"ok": True, "documents": document_formats if available else [], "pdf": pdf,
            "image_ocr": image_formats, "ocr": ocr["ocr"], "langs": ocr["langs"]}


def preflight_zip(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 20000 or sum(x.file_size for x in entries) > MAX_OUTPUT:
                raise ConversionError("O documento compactado excede o limite de recursos.")
            if any(x.flag_bits & 1 for x in entries):
                raise ConversionError("Documentos protegidos por senha não são suportados.")
    except zipfile.BadZipFile:
        raise ConversionError("O documento está danificado ou possui formato incorreto.") from None


def write_profile(directory: Path) -> None:
    user = directory / "profile/user"
    user.mkdir(parents=True)
    (directory / "home").mkdir()
    (directory / "tmp").mkdir()
    (user / "registrymodifications.xcu").write_text(
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<oor:items xmlns:oor="http://openoffice.org/2001/registry">'
        '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
        '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop>'
        '<prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop>'
        '<prop oor:name="DisableActiveContent" oor:op="fuse"><value>true</value></prop></item>'
        '<item oor:path="/org.openoffice.Office.Java/VirtualMachine">'
        '<prop oor:name="Enable" oor:op="fuse"><value>false</value></prop></item>'
        '<item oor:path="/org.openoffice.Office.Writer/Content/Update">'
        '<prop oor:name="Link" oor:op="fuse"><value>0</value></prop>'
        '<prop oor:name="Field" oor:op="fuse"><value>false</value></prop>'
        '<prop oor:name="Chart" oor:op="fuse"><value>false</value></prop></item>'
        '</oor:items>', encoding="utf-8")


def office_command(root: Path, work: Path, extension: str, target: str) -> list[str]:
    # Only the installation, system libraries/fonts and this job's directory are
    # visible. There is no network or access to the user's files/profiles/sockets.
    isolation = office_isolation_tool()
    if isolation is None:
        raise ConversionError("O conversor local de documentos não está disponível.")
    command = [str(isolation), "--die-with-parent", "--unshare-net",
               "--ro-bind", "/usr", "/usr",
               "--ro-bind", "/bin", "/bin",
               "--ro-bind", "/lib", "/lib",
               "--ro-bind-try", "/lib64", "/lib64",
               "--ro-bind-try", "/etc/ld.so.cache", "/etc/ld.so.cache",
               "--ro-bind-try", "/etc/fonts", "/etc/fonts",
               "--ro-bind-try", "/var/cache/fontconfig", "/var/cache/fontconfig",
               "--dev", "/dev", "--proc", "/proc", "--tmpfs", "/tmp",
               "--ro-bind", str(root), "/opt/maracuja-libreoffice",
               "--bind", str(work), "/work", "--chdir", "/work"]
    # AppRun's shared-library path must not leak into the private Python runtime.
    # Supply the portable Writer dependencies only after entering its sandbox,
    # where the host's AppImage/extraction path is deliberately not exposed.
    libraries = runtime().root / "lib"
    if libraries.is_dir():
        command.extend(["--ro-bind", str(libraries.resolve()), "/opt/maracuja-libs",
                        "--setenv", "LD_LIBRARY_PATH", "/opt/maracuja-libs"])
    fonts = runtime().root / "fonts"
    if (fonts / "fonts.conf").is_file():
        command.extend(["--ro-bind", str(fonts.resolve()), "/opt/maracuja-fonts",
                        "--setenv", "FONTCONFIG_FILE", "/opt/maracuja-fonts/fonts.conf",
                        "--setenv", "XDG_CACHE_HOME", "/work/home/.cache"])
    command.extend(["/opt/maracuja-libreoffice/program/soffice",
               "-env:UserInstallation=file:///work/profile", "--headless", "--nologo",
               "--nodefault", "--nofirststartwizard", "--norestore"])
    if extension == "txt":
        command.append("--infilter=Text (encoded):UTF8,LF,Liberation Sans,en-US")
    elif extension in {"md", "markdown"}:
        command.append("--infilter=Markdown")
    command.extend(["--convert-to", FILTERS[target], "--outdir", "/work/out",
                    "/work/source." + extension])
    return command


def prepare_windows_source(path: Path, extension: str) -> None:
    """Reject active/external Office resources before invoking Windows Writer.

    This validation is deliberately conservative. It is not an OS sandbox and
    does not claim Bubblewrap's process/filesystem/network isolation.
    """
    message = ("O documento contém conteúdo ativo ou recursos externos. "
               "No Windows, salve uma cópia autocontida antes de converter.")

    def external(reference: str) -> bool:
        from urllib.parse import unquote, urlsplit
        value = unquote(reference.strip())
        if not value or value.startswith("#"):
            return False
        return bool(urlsplit(value).scheme or urlsplit(value).netloc
                    or value.startswith(("/", "\\")) or "\\" in value
                    or ".." in Path(value).parts)

    if extension in {"docx", "odt"}:
        preflight_zip(path)
        with zipfile.ZipFile(path) as archive:
            for entry in archive.infolist():
                name = entry.filename.lower()
                if any(part in name.split("/") for part in ["activex", "scripts", "basic", "embeddings"]):
                    raise ConversionError(message)
                if "vbaproject" in name or name.endswith((".html", ".htm", ".mht", ".mhtml", ".svg", ".svgz")):
                    raise ConversionError(message)
                if not name.endswith((".xml", ".rels")):
                    continue
                if entry.file_size > 32 * 1024 * 1024:
                    raise ConversionError("O conteúdo do documento excede o limite de recursos.")
                data = archive.read(entry)
                if re.search(br"<!\s*(DOCTYPE|ENTITY)\b", data.replace(b"\x00", b""), re.IGNORECASE):
                    raise ConversionError(message)
                try:
                    tree = ET.fromstring(data)
                except ET.ParseError:
                    raise ConversionError("O documento contém conteúdo XML inválido.") from None
                field_instructions = []
                for element in tree.iter():
                    tag = element.tag.rsplit("}", 1)[-1].lower()
                    if tag in {"oleobject", "object", "altchunk", "ddelink", "ddelinks", "externaldata"}:
                        raise ConversionError(message)
                    if tag == "instrtext":
                        field_instructions.append(element.text or "")
                    for key, value in element.attrib.items():
                        attribute = key.rsplit("}", 1)[-1].lower()
                        if attribute == "targetmode" and value.lower() == "external":
                            raise ConversionError(message)
                        if attribute in {"href", "src", "base"} and external(value):
                            raise ConversionError(message)
                        # OOXML package relationship targets may contain ../ to
                        # reach another part inside the ZIP; URI schemes may not.
                        if attribute == "target":
                            from urllib.parse import urlsplit
                            parsed = urlsplit(value)
                            if parsed.scheme or parsed.netloc or value.startswith(("/", "\\")):
                                raise ConversionError(message)
                        if attribute == "instr":
                            field_instructions.append(value)
                if re.search(r"\b(INCLUDEPICTURE|INCLUDETEXT|DDE|DDEAUTO|LINK|DATABASE)\b",
                             "".join(field_instructions), re.IGNORECASE):
                    raise ConversionError(message)
    elif extension == "txt":
        path.write_text(decode_text_document(path), encoding="utf-8")
    elif extension == "rtf":
        data = path.read_bytes()
        if not data.lstrip().startswith(b"{\\rtf"):
            raise ConversionError("O arquivo não contém um documento RTF válido.")
        if re.search(br"\\(field|fldinst|object|objdata|filetbl|datastore|htmltag|fromhtml)\b",
                     data, re.IGNORECASE):
            raise ConversionError(message)
    elif extension in {"md", "markdown"}:
        text = decode_text_document(path)
        if re.search(r"!\s*\[|<|^\s*\[[^\]]+\]:", text, re.MULTILINE):
            raise ConversionError(message)
        # Hyperlink labels remain; Writer never receives their destinations.
        text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
        path.write_text(text, encoding="utf-8")
    elif extension in {"html", "htm"}:
        from html.parser import HTMLParser

        class LocalHTML(HTMLParser):
            tags = {"html", "head", "body", "title", "p", "div", "span", "br", "hr",
                    "h1", "h2", "h3", "h4", "h5", "h6", "b", "strong", "i", "em", "u",
                    "s", "del", "sup", "sub", "pre", "code", "blockquote", "ol", "ul", "li",
                    "table", "thead", "tbody", "tfoot", "tr", "td", "th", "caption", "dl",
                    "dt", "dd", "a", "img"}
            blocked = {"script", "style", "object", "iframe", "svg", "math", "template"}

            def __init__(self):
                super().__init__(convert_charrefs=True)
                self.output = ['<!doctype html><html><head><meta charset="UTF-8"></head><body>']
                self.ignored = []
                self.image_bytes = 0

            def handle_starttag(self, tag, attributes):
                if self.ignored:
                    if tag in self.blocked:
                        self.ignored.append(tag)
                    return
                if tag in self.blocked:
                    self.ignored.append(tag)
                    return
                if tag not in self.tags or tag in {"html", "head", "body", "title"}:
                    return
                allowed = []
                for key, value in attributes:
                    if value is None:
                        continue
                    if key in {"title", "alt"}:
                        allowed.append((key, value))
                    elif key in {"width", "height", "colspan", "rowspan"} and re.fullmatch(r"\d{1,5}", value):
                        allowed.append((key, value))
                    elif key == "href" and tag == "a" and value.startswith("#"):
                        allowed.append((key, value))
                    elif key == "src" and tag == "img":
                        matched = re.fullmatch(r"data:image/(png|jpeg|gif|webp);base64,([A-Za-z0-9+/=\s]+)", value,
                                               re.IGNORECASE)
                        if matched is None:
                            raise ConversionError(message)
                        encoded = re.sub(r"\s", "", matched[2])
                        if len(encoded) > 24 * 1024 * 1024:
                            raise ConversionError("Uma imagem do documento excede o limite de tamanho.")
                        try:
                            decoded = base64.b64decode(encoded, validate=True)
                        except ValueError:
                            raise ConversionError("O documento contém uma imagem inválida.") from None
                        # A data:image/png label can hide SVG/HTML. Decode and
                        # re-encode real raster pixels before giving them to LO.
                        from PIL import Image
                        try:
                            with Image.open(io.BytesIO(decoded)) as original:
                                if original.format not in {"PNG", "JPEG", "GIF", "WEBP"} or original.width * original.height > 16_000_000:
                                    raise ConversionError("O documento contém uma imagem não suportada.")
                                pixels = original.convert("RGBA")
                                try:
                                    with io.BytesIO() as buffer:
                                        pixels.save(buffer, format="PNG")
                                        raster = buffer.getvalue()
                                finally:
                                    pixels.close()
                        except ConversionError:
                            raise
                        except Exception:
                            raise ConversionError("O documento contém uma imagem inválida.") from None
                        self.image_bytes += len(raster)
                        if self.image_bytes > 96 * 1024 * 1024:
                            raise ConversionError("As imagens do documento excedem o limite de tamanho.")
                        allowed.append((key, "data:image/png;base64," + base64.b64encode(raster).decode("ascii")))
                self.output.append("<" + tag + "".join(" " + key + '=\"' + html.escape(value, quote=True)
                                                         + '\"' for key, value in allowed) + ">")

            def handle_startendtag(self, tag, attributes):
                self.handle_starttag(tag, attributes)
                if tag not in {"br", "hr", "img"}:
                    self.handle_endtag(tag)

            def handle_endtag(self, tag):
                if self.ignored:
                    if tag == self.ignored[-1]:
                        self.ignored.pop()
                    return
                if tag in self.tags - {"html", "head", "body", "title", "br", "hr", "img"}:
                    self.output.append("</" + tag + ">")

            def handle_data(self, data):
                if not self.ignored:
                    self.output.append(html.escape(data))

        sanitizer = LocalHTML()
        sanitizer.feed(decode_text_document(path))
        sanitizer.close()
        path.write_text("".join(sanitizer.output) + "</body></html>", encoding="utf-8")


def decode_text_document(path: Path) -> str:
    data = path.read_bytes()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def windows_office_command(executable: Path, work: Path, extension: str, target: str) -> list[str]:
    filters = {"txt": "Text (encoded):UTF8,LF,Liberation Sans,en-US", "md": "Markdown",
               "markdown": "Markdown", "rtf": "Rich Text Format", "html": "HTML (StarWriter)",
               "htm": "HTML (StarWriter)", "docx": "Office Open XML Text", "odt": "writer8"}
    return [str(executable), "-env:UserInstallation=" + (work / "profile").as_uri(),
            "--headless", "--nologo", "--nodefault", "--nofirststartwizard", "--norestore",
            "--infilter=" + filters[extension], "--convert-to", FILTERS[target],
            "--outdir", str(work / "out"), str(work / ("source." + extension))]


def embedded_html(path: Path, directory: Path) -> str:
    text = path.read_text(encoding="utf-8", errors="replace")
    if len(text.encode("utf-8")) > MAX_OUTPUT:
        raise ConversionError("O documento convertido excede o limite de tamanho.")
    total = 0

    def data_uri(reference: str) -> str:
        nonlocal total
        from urllib.parse import unquote, urlparse
        value = html.unescape(reference.strip())
        parsed = urlparse(value)
        if parsed.scheme == "data":
            return value
        if parsed.scheme or parsed.netloc:
            return ""  # Never fetch remote or file:// references.
        asset = (directory / unquote(parsed.path)).resolve()
        if not asset.is_relative_to(directory.resolve()) or not asset.is_file():
            return ""
        size = asset.stat().st_size
        if size > 16 * 1024 * 1024 or total + size > 96 * 1024 * 1024:
            raise ConversionError("As imagens do documento excedem o limite de tamanho.")
        mime = mimetypes.guess_type(asset.name)[0] or "application/octet-stream"
        if not mime.startswith("image/"):
            return ""
        total += size
        return "data:" + mime + ";base64," + base64.b64encode(asset.read_bytes()).decode("ascii")

    # LibreOffice produces quoted HTML attributes and quoted/unquoted CSS URLs.
    text = re.sub(r"\b(src|background)\s*=\s*([\"'])(.*?)\2",
                  lambda m: m[1] + '="' + html.escape(data_uri(m[3]), quote=True) + '"',
                  text, flags=re.IGNORECASE | re.DOTALL)
    text = re.sub(r"url\(\s*([\"']?)(.*?)\1\s*\)",
                  lambda m: 'url("' + data_uri(m[2]) + '")', text,
                  flags=re.IGNORECASE | re.DOTALL)
    # Do not leave relative stylesheets or executable content in an exported file.
    text = re.sub(r"<(script|iframe|object)\b[^>]*>.*?</\1\s*>", "", text,
                  flags=re.IGNORECASE | re.DOTALL)
    text = re.sub(r"<(link|embed|base)\b[^>]*>", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+on[a-z]+\s*=\s*([\"']).*?\1", "", text,
                  flags=re.IGNORECASE | re.DOTALL)
    text = re.sub(r"(href\s*=\s*[\"'])\s*javascript:[^\"']*", r"\1#", text,
                  flags=re.IGNORECASE)
    text = re.sub(r"@import\s+[^;]+;", "", text, flags=re.IGNORECASE)
    text = re.sub(r"<meta\b[^>]*http-equiv\s*=\s*[\"']?refresh[^>]*>", "", text,
                  flags=re.IGNORECASE)
    return text


@contextmanager
def document_workspace(parent: Path):
    directory = os.path.abspath(parent)
    if os.name == "nt" and not directory.startswith("\\\\?\\"):
        # Writer's extension registry can exceed MAX_PATH even for an ordinary
        # output folder. Keep extended paths for TemporaryDirectory's cleanup,
        # including its handling of read-only Windows profile directories.
        directory = "\\\\?\\UNC\\" + directory[2:] if directory.startswith("\\\\") else "\\\\?\\" + directory
    temporary = tempfile.TemporaryDirectory(prefix=".maracuja-documento-", dir=directory)
    work = temporary.name
    # Writer expects conventional paths and file:// profile URLs.
    if os.name == "nt":
        if work.startswith("\\\\?\\UNC\\"):
            work = "\\\\" + work[8:]
        elif work.startswith("\\\\?\\"):
            work = work[4:]
    try:
        yield Path(work)
    except BaseException:
        try:
            temporary.cleanup()
        except OSError:
            pass  # A cleanup failure must not replace the conversion error.
        raise
    else:
        temporary.cleanup()


def convert_document(source: Path, target: str, output: Path, soffice: str) -> None:
    root = office_root(soffice)
    if root is None or (not runtime().windows and office_isolation_tool() is None):
        raise ConversionError("O conversor local de documentos não está disponível.")
    extension = source.suffix.lower().lstrip(".")
    if (target == "md" or extension in {"md", "markdown"}) and not platform_module().markdown_filter(root):
        raise ConversionError("Este LibreOffice não oferece conversão de Markdown. Atualize o componente de documentos.")
    if extension in {"docx", "odt"}:
        preflight_zip(source)
    with document_workspace(output.parent) as work:
        write_profile(work)
        (work / "out").mkdir()
        # Do not expose the source directory or any unrelated user file to Writer.
        shutil.copyfile(source, work / ("source." + extension))
        if runtime().windows:
            prepare_windows_source(work / ("source." + extension), extension)
            executable = runtime().tool("soffice", soffice)
            if executable is None:
                raise ConversionError("O conversor local de documentos não está disponível.")
            command = windows_office_command(executable, work, extension, target)
            environment = dict(os.environ)
            roaming, local = work / "home/AppData/Roaming", work / "home/AppData/Local"
            roaming.mkdir(parents=True)
            local.mkdir(parents=True)
            environment.update({"HOME": str(work / "home"), "USERPROFILE": str(work / "home"),
                                "APPDATA": str(roaming), "LOCALAPPDATA": str(local),
                                "TEMP": str(work / "tmp"), "TMP": str(work / "tmp")})
            system = Path(environment.get("SystemRoot", r"C:\Windows"))
            environment["PATH"] = os.pathsep.join([str(executable.parent), str(runtime().root / "tools"),
                                                    str(system / "System32")])
        else:
            command = office_command(root, work, extension, target)
            environment = {"PATH": "/usr/bin:/bin", "LANG": "C.UTF-8", "LC_ALL": "C.UTF-8",
                           "HOME": "/work/home", "TMPDIR": "/work/tmp", "SAL_USE_VCLPLUGIN": "svp"}
        try:
            result = subprocess.run(command,
                                    stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                                    stderr=subprocess.DEVNULL, env=environment, timeout=120, check=False,
                                    **runtime().subprocess_options())
        except subprocess.TimeoutExpired:
            raise ConversionError("A conversão do documento excedeu o limite de tempo.") from None
        product = work / "out" / ("source." + target)
        if result.returncode != 0 or not product.is_file() or not 0 < product.stat().st_size <= MAX_OUTPUT:
            raise ConversionError("Não foi possível ler ou converter este documento.")
        if target == "html":
            output.write_text(embedded_html(product, work / "out"), encoding="utf-8")
        else:
            shutil.copyfile(product, output)


def open_pdf(source: Path):
    from pypdf import PdfReader
    try:
        reader = PdfReader(str(source), strict=False)
        if reader.is_encrypted:
            raise ConversionError("PDFs protegidos por senha não são suportados.")
        if not 0 < len(reader.pages) <= MAX_PAGES:
            raise ConversionError("Esta conversão aceita PDFs de até 200 páginas.")
        for page in reader.pages:
            width, height = float(page.mediabox.width), float(page.mediabox.height)
            if not all(math.isfinite(x) and 0 < x <= 144000 for x in [width, height]):
                raise ConversionError("O PDF contém dimensões de página inválidas.")
        return reader
    except ConversionError:
        raise
    except Exception:
        raise ConversionError("Não foi possível ler este PDF.") from None


def ocr_pdf_page(source: Path, page_number: int, output: Path, engine, cropbox) -> str:
    from PIL import Image
    width, height = float(cropbox.width), float(cropbox.height)
    if not all(math.isfinite(side) and side > 0 for side in (width, height)):
        raise ConversionError("Uma página para OCR possui dimensões inválidas.")
    # Render at 300 DPI, reducing only to respect the memory limit. Forcing
    # every page to 4096px also enlarges small/cropped pages unnecessarily.
    dpi = min(300.0, 4095.0 * 72.0 / max(width, height))
    with tempfile.TemporaryDirectory(prefix=".maracuja-ocr-", dir=output.parent) as temporary:
        prefix = Path(temporary) / "page"
        try:
            result = subprocess.run([required_tool("pdftoppm"), "-f", str(page_number), "-l", str(page_number),
                                     "-singlefile", "-r", str(dpi), "-cropbox", "-gray", "-png",
                                     str(source), str(prefix)], stdin=subprocess.DEVNULL,
                                    stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                                    timeout=45, check=False, **runtime().subprocess_options())
        except subprocess.TimeoutExpired:
            raise ConversionError("A preparação de uma página para OCR excedeu o limite de tempo.") from None
        image_path = prefix.with_suffix(".png")
        if result.returncode or not image_path.is_file() or image_path.stat().st_size > MAX_SOURCE:
            raise ConversionError("Não foi possível preparar uma página do PDF para OCR.")
        with Image.open(image_path) as image:
            if image.width < 1 or image.height < 1 or max(image.size) > 4096:
                raise ConversionError("Uma página para OCR possui dimensões inválidas.")
            if image.width * image.height > ocr_module().MAX_PIXELS:
                image.thumbnail((4000, 4000), Image.Resampling.LANCZOS)
            try:
                return clean_text(engine.recognize(image, round(dpi)))
            except ocr_module().OCRError as error:
                raise ConversionError(str(error)) from None


def pdf_page_text(page, source: Path, index: int, output: Path, remaining: int) -> str:
    replacement_text = False

    def inspect_marked_content(operator, operands, _current_matrix, _text_matrix):
        nonlocal replacement_text
        if operator == b"BDC" and len(operands) >= 2:
            properties = operands[1]
            # Named property lists can live inside nested form resources, which
            # this visitor does not expose. Let Poppler resolve those as well.
            if isinstance(properties, str) or (isinstance(properties, dict) and "/ActualText" in properties):
                replacement_text = True

    text = clean_text(page.extract_text(visitor_operand_before=inspect_marked_content) or "")
    if not replacement_text:
        return text
    executable = required_tool("pdftotext")
    # Shaped fonts can represent one Unicode character with several glyphs.
    # Poppler honors PDF /ActualText spans; pypdf's extraction currently does
    # not. Raw content order also avoids inserting spaces inside those spans.
    with tempfile.TemporaryDirectory(prefix=".maracuja-texto-", dir=output.parent) as temporary:
        text_path = Path(temporary) / "page.txt"
        try:
            result = subprocess.run([str(executable), "-raw", "-enc", "UTF-8", "-nopgbrk",
                                     "-f", str(index), "-l", str(index), str(source), str(text_path)],
                                    stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                                    stderr=subprocess.DEVNULL, timeout=30, check=False,
                                    **runtime().subprocess_options())
        except subprocess.TimeoutExpired:
            raise ConversionError("A extração do texto do PDF excedeu o limite de tempo.") from None
        if result.returncode or not text_path.is_file():
            raise ConversionError("Não foi possível extrair o texto de uma página do PDF.")
        if text_path.stat().st_size > remaining:
            raise ConversionError("O texto do PDF excede o limite de recursos.")
        return clean_text(text_path.read_text(encoding="utf-8"))


def pdf_text(reader, source: Path, output: Path, tessdata: str = "", progress=emit_progress) -> list[str]:
    pages = []
    length = 0
    engine = None
    try:
        ocr_available = ocr_module().capabilities(tessdata, _RUNTIME_ROOT)["ocr"]
    except Exception:
        ocr_available = False
    total = len(reader.pages)
    try:
        for index, page in enumerate(reader.pages, 1):
            progress(index - 1, total, "text")
            text = pdf_page_text(page, source, index, output, MAX_TEXT - length)
            phase = "text"
            # Never OCR/rewrite a page that already contains selectable text.
            if not text.strip() and ocr_available:
                phase = "ocr"
                progress(index - 1, total, phase)
                if engine is None:
                    try:
                        engine = ocr_module().Engine(tessdata, _RUNTIME_ROOT)
                    except ocr_module().OCRError as error:
                        raise ConversionError(str(error)) from None
                text = ocr_pdf_page(source, index, output, engine, page.cropbox)
            length += len(text.encode("utf-8"))
            if length > MAX_TEXT:
                raise ConversionError("O texto do PDF excede o limite de recursos.")
            pages.append(text)
            progress(index, total, phase)
    finally:
        if engine is not None:
            engine.close()
    if not any(page.strip() for page in pages):
        if not ocr_available:
            raise ConversionError("O PDF não possui texto selecionável e o OCR local não está disponível.")
        raise ConversionError("Nenhum texto legível foi encontrado no PDF, mesmo após o OCR.")
    return pages


def plain_paragraphs(text: str) -> list[str]:
    return [x.strip() for x in re.split(r"\n\s*\n", text) if x.strip()]


def text_html(pages: list[str]) -> str:
    sections = []
    for index, text in enumerate(pages, 1):
        paragraphs = "".join('<p style="white-space:pre-wrap">' + html.escape(p) + "</p>"
                             for p in plain_paragraphs(text))
        sections.append('<section><h2>Página ' + str(index) + "</h2>" + paragraphs + "</section>")
    return ('<!doctype html><html lang="pt-BR"><head><meta charset="UTF-8">'
            '<meta name="viewport" content="width=device-width, initial-scale=1">'
            '<title>Documento convertido</title></head><body>' + "".join(sections) + "</body></html>")


def write_epub(pages: list[str], output: Path) -> None:
    from xml.sax.saxutils import escape
    from uuid import uuid4
    from datetime import datetime, timezone
    identifier = "urn:uuid:" + str(uuid4())
    modified = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    chapters = []
    for index, text in enumerate(pages, 1):
        body = "".join('<p style="white-space:pre-wrap">' + escape(p) + '</p>'
                       for p in plain_paragraphs(text))
        chapters.append('<?xml version="1.0" encoding="utf-8"?>'
                        '<html xmlns="http://www.w3.org/1999/xhtml" lang="pt-BR"><head>'
                        '<title>Página ' + str(index) + '</title></head><body><h1>Página '
                        + str(index) + '</h1>' + body + '</body></html>')
    manifest = ''.join('<item id="p' + str(i) + '" href="page-' + str(i)
                       + '.xhtml" media-type="application/xhtml+xml"/>'
                       for i in range(1, len(chapters) + 1))
    spine = ''.join('<itemref idref="p' + str(i) + '"/>' for i in range(1, len(chapters) + 1))
    nav = ''.join('<li><a href="page-' + str(i) + '.xhtml">Página ' + str(i) + '</a></li>'
                  for i in range(1, len(chapters) + 1))
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/container.xml", '<?xml version="1.0"?><container version="1.0" '
                         'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile '
                         'full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                         '</rootfiles></container>')
        archive.writestr("OEBPS/content.opf", '<?xml version="1.0" encoding="utf-8"?>'
                         '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book">'
                         '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="book">'
                         + identifier + '</dc:identifier><dc:title>Documento convertido</dc:title>'
                         '<dc:language>pt-BR</dc:language><meta property="dcterms:modified">' + modified
                         + '</meta></metadata><manifest><item id="nav" href="nav.xhtml" '
                         'media-type="application/xhtml+xml" properties="nav"/>' + manifest
                         + '</manifest><spine>' + spine + '</spine></package>')
        archive.writestr("OEBPS/nav.xhtml", '<?xml version="1.0" encoding="utf-8"?>'
                         '<html xmlns="http://www.w3.org/1999/xhtml" '
                         'xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Sumário</title>'
                         '</head><body><nav epub:type="toc"><h1>Sumário</h1><ol>' + nav
                         + '</ol></nav></body></html>')
        for i, chapter in enumerate(chapters, 1):
            archive.writestr("OEBPS/page-" + str(i) + ".xhtml", chapter)


def pdf_presentation(source: Path, reader, output: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333333), Inches(7.5)
    total_image_bytes = 0
    with tempfile.TemporaryDirectory(prefix=".maracuja-slides-", dir=output.parent) as temporary:
        directory = Path(temporary)
        for index in range(1, len(reader.pages) + 1):
            prefix = directory / "page"
            result = subprocess.run([required_tool("pdftoppm"), "-f", str(index), "-l", str(index),
                                     "-singlefile", "-scale-to", "1920", "-png", str(source), str(prefix)],
                                    stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                                    stderr=subprocess.DEVNULL, timeout=30, check=False,
                                    **runtime().subprocess_options())
            image = prefix.with_suffix(".png")
            if result.returncode or not image.is_file() or image.stat().st_size > 16 * 1024 * 1024:
                raise ConversionError("Não foi possível renderizar uma página do PDF.")
            total_image_bytes += image.stat().st_size
            if total_image_bytes > 128 * 1024 * 1024:
                raise ConversionError("As páginas do PDF excedem o limite da apresentação.")
            with Image.open(image) as frame:
                width, height = frame.size
            if width <= 0 or height <= 0 or width * height > 4 * 1024 * 1024:
                raise ConversionError("Uma página do PDF excede o limite de imagem.")
            scale = min(presentation.slide_width / width, presentation.slide_height / height)
            picture_width, picture_height = int(width * scale), int(height * scale)
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_picture(str(image), (presentation.slide_width - picture_width) // 2,
                                    (presentation.slide_height - picture_height) // 2,
                                    width=picture_width, height=picture_height)
            image.unlink()
        presentation.save(str(output))


def pdf_tables(source: Path, output: Path) -> None:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    workbook = Workbook()
    workbook.remove(workbook.active)
    table_count = cell_count = text_size = 0
    with pdfplumber.open(source) as document:
        for page_index, page in enumerate(document.pages, 1):
            try:
                for table_index, table in enumerate(page.extract_tables(), 1):
                    if not any(any(str(cell or "").strip() for cell in row) for row in table):
                        continue
                    table_count += 1
                    if table_count > 500:
                        raise ConversionError("O PDF contém tabelas demais para esta conversão.")
                    worksheet = workbook.create_sheet("P" + str(page_index) + " T" + str(table_index))
                    for row_index, row in enumerate(table, 1):
                        cell_count += len(row)
                        if cell_count > 200000 or len(row) > 16384:
                            raise ConversionError("As tabelas do PDF excedem o limite de recursos.")
                        for column, cell in enumerate(row, 1):
                            value = ILLEGAL_CHARACTERS_RE.sub("", str(cell or ""))[:32767]
                            text_size += len(value)
                            if text_size > MAX_TEXT:
                                raise ConversionError("As tabelas do PDF excedem o limite de texto.")
                            target = worksheet.cell(row_index, column, value)
                            # Extracted PDF strings are data, never spreadsheet formulae.
                            target.data_type = "s"
                    worksheet.freeze_panes = "A2"
            finally:
                page.close()
    if not table_count:
        raise ConversionError("Nenhuma tabela foi encontrada neste PDF. PDFs digitalizados precisam de OCR.")
    workbook.save(str(output))


def write_plain_document(pages: list[str], target: str, output: Path) -> None:
    if target == "txt":
        output.write_text("\n\f\n".join(page.rstrip() for page in pages) + "\n", encoding="utf-8")
    elif target == "docx":
        from docx import Document
        from docx.shared import Pt
        document = Document()
        document.styles["Normal"].font.name = "Liberation Sans"
        document.styles["Normal"].font.size = Pt(11)
        for index, text in enumerate(pages):
            if index:
                document.add_page_break()
            for paragraph in plain_paragraphs(text):
                document.add_paragraph(clean_text(paragraph))
        document.save(str(output))
    else:
        raise ConversionError("Este formato de texto não está disponível.")


def convert_image_ocr(source: Path, target: str, output: Path,
                      tessdata: str = "", progress=emit_progress) -> None:
    from PIL import Image
    progress(0, 1, "ocr")
    try:
        with Image.open(source) as image:
            # Only Qt's normalized, single-frame PNG is accepted by this route.
            if image.format != "PNG" or getattr(image, "n_frames", 1) != 1:
                raise ConversionError("A imagem para OCR precisa ser um PNG normalizado.")
            if image.width * image.height > ocr_module().MAX_PIXELS:
                raise ConversionError("A imagem para OCR excede o limite de 16 milhões de pixels.")
            with ocr_module().Engine(tessdata, _RUNTIME_ROOT) as engine:
                text = clean_text(engine.recognize(image, 300))
    except ocr_module().OCRError as error:
        raise ConversionError(str(error)) from None
    if not text.strip():
        raise ConversionError("Nenhum texto legível foi encontrado na imagem.")
    if len(text.encode("utf-8")) > MAX_TEXT:
        raise ConversionError("O texto reconhecido excede o limite de tamanho.")
    progress(1, 1, "ocr")
    write_plain_document([text], target, output)


def convert_pdf(source: Path, target: str, output: Path,
                tessdata: str = "", progress=emit_progress) -> None:
    reader = open_pdf(source)
    if target == "pptx":
        pdf_presentation(source, reader, output)
        return
    if target == "xlsx":
        pdf_tables(source, output)
        return
    pages = pdf_text(reader, source, output, tessdata, progress)
    if target in {"txt", "docx"}:
        write_plain_document(pages, target, output)
    elif target == "html":
        output.write_text(text_html(pages), encoding="utf-8")
    elif target == "md":
        # Plain PDF text must not silently turn into Markdown syntax/HTML.
        escaped = [re.sub(r"([\\`*_{}\[\]<>#+.!|~-])", r"\\\1", text) for text in pages]
        output.write_text("\n\n---\n\n".join(escaped), encoding="utf-8")
    elif target == "epub":
        write_epub(pages, output)
    else:
        raise ConversionError("Este formato de PDF não está disponível.")


def main() -> int:
    global _RUNTIME_ROOT, _RUNTIME
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--capabilities", action="store_true")
    parser.add_argument("--input")
    parser.add_argument("--output")
    parser.add_argument("--format")
    parser.add_argument("--soffice", default="")
    parser.add_argument("--tessdata", default="")
    parser.add_argument("--image-ocr", action="store_true")
    parser.add_argument("--runtime-root", default="")
    arguments = parser.parse_args()
    _RUNTIME_ROOT = arguments.runtime_root
    _RUNTIME = None
    if arguments.capabilities:
        print(json.dumps(capabilities(arguments.soffice, arguments.tessdata)), flush=True)
        return 0
    output = None
    cleanup_output = False
    try:
        if not all([arguments.input, arguments.output, arguments.format]):
            raise ConversionError("Parâmetros de conversão inválidos.")
        source, output = Path(arguments.input).resolve(), Path(arguments.output).absolute()
        target = arguments.format.lower()
        extension = source.suffix.lower().lstrip(".")
        if not source.is_file() or not 0 <= source.stat().st_size <= MAX_SOURCE:
            raise ConversionError("Documento inválido ou maior que 64 MiB.")
        if output.exists() or output.is_symlink() or source == output.resolve() or not output.parent.is_dir():
            raise ConversionError("O destino temporário do documento é inválido.")
        cleanup_output = True
        supported = capabilities(arguments.soffice, arguments.tessdata)
        is_pdf = extension == "pdf"
        if extension not in DOCUMENT_EXTENSIONS and not is_pdf:
            with source.open("rb") as stream:
                is_pdf = stream.read(5) == b"%PDF-"
        if arguments.image_ocr and target in supported["image_ocr"]:
            convert_image_ocr(source, target, output, arguments.tessdata)
        elif arguments.image_ocr:
            raise ConversionError("Este formato de reconhecimento de texto não está disponível.")
        elif is_pdf and target in supported["pdf"]:
            convert_pdf(source, target, output, arguments.tessdata)
        elif extension in DOCUMENT_EXTENSIONS and target in supported["documents"]:
            convert_document(source, target, output, arguments.soffice)
        else:
            raise ConversionError("Esta combinação de documentos não está disponível.")
        if not output.is_file() or not 0 < output.stat().st_size <= MAX_OUTPUT:
            raise ConversionError("A conversão não produziu um documento válido dentro do limite de tamanho.")
        print(json.dumps({"ok": True}), flush=True)
        return 0
    except BaseException as error:
        if output is not None and cleanup_output:
            try:
                output.unlink(missing_ok=True)
            except OSError:
                pass
        message = str(error) if isinstance(error, ConversionError) else "Não foi possível ler ou converter este documento."
        print(json.dumps({"ok": False, "error": message}, ensure_ascii=True), flush=True)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
