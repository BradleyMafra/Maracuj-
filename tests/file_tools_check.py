#!/usr/bin/env python3
"""Generated fixtures only: lossless metadata changes, copy safety and compression."""
from __future__ import annotations
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest import mock
import warnings
import zipfile
import zlib
import xml.etree.ElementTree as ET
import sys
import os
import subprocess

from PIL import Image
from PIL.PngImagePlugin import PngInfo
from pypdf import PdfReader, PdfWriter

HELPER = Path(sys.argv.pop(1)).resolve() if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / 'helpers/file_tools.py'
spec = importlib.util.spec_from_file_location('maracuja_file_tools_tested', HELPER)
tools = importlib.util.module_from_spec(spec)
spec.loader.exec_module(tools)


class FileToolsCheck(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='maracuja-tools-test-')
        self.root = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def edit(self, path, options, action='metadata', name='result'):
        destination = self.root / (name + path.suffix)
        original = hashlib.sha256(path.read_bytes()).hexdigest()
        result = tools.execute({'action': action, 'inputs': [str(path)], 'output': str(destination), 'options': options})
        self.assertTrue(result['ok'])
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), original)
        self.assertTrue(destination.is_file())
        return destination, result

    def jpg(self):
        path = self.root / 'photo.jpg'
        exif = Image.Exif()
        exif[270], exif[315], exif[271], exif[274] = 'Before', 'Private Author', 'Private Camera', 6
        Image.new('RGB', (80, 60), '#235678').save(path, quality=95, exif=exif)
        return path

    def test_jpeg_edit_preserves_encoded_scan(self):
        source = self.jpg()
        output, _ = self.edit(source, {'changes': {'exif:270': 'Depois'}, 'remove': ['exif:315']})
        self.assertEqual(tools.jpeg_segments(source.read_bytes())[1], tools.jpeg_segments(output.read_bytes())[1])
        with Image.open(output) as image:
            self.assertEqual(image.getexif()[270], 'Depois')
            self.assertNotIn(315, image.getexif())
            self.assertEqual(image.getexif()[274], 6)
        self.assertTrue(next(r for r in tools.inspect(output)['metadata'] if r['key'] == 'image:width')['removable'] is False)

    def test_jpeg_remove_all_preserves_orientation_and_pixels(self):
        source = self.jpg()
        output, _ = self.edit(source, {'removeAll': True})
        self.assertEqual(tools.jpeg_segments(source.read_bytes())[1], tools.jpeg_segments(output.read_bytes())[1])
        with Image.open(output) as image:
            self.assertEqual(dict(image.getexif()), {274: 6})

    def test_png_text_and_exif_edit_without_idat_change(self):
        source = self.root / 'photo.png'
        text = PngInfo()
        text.add_text('Author', 'Private Author')
        text.add_text('Secret', 'Private Other')
        Image.new('RGBA', (81, 61), '#8723a870').save(source, pnginfo=text)
        output, _ = self.edit(source, {'changes': {'text:Author': 'Novo Autor'}})
        old = [data for kind, data in tools.png_chunks(source.read_bytes()) if kind == b'IDAT']
        new = [data for kind, data in tools.png_chunks(output.read_bytes()) if kind == b'IDAT']
        self.assertEqual(old, new)
        with Image.open(output) as image:
            self.assertEqual(image.text['Author'], 'Novo Autor')
        cleaned, _ = self.edit(source, {'removeAll': True}, name='cleaned')
        with Image.open(cleaned) as image:
            self.assertEqual(image.text, {})
            self.assertEqual(image.getpixel((5, 5)), (135, 35, 168, 112))

    def test_webp_metadata_does_not_change_pixels(self):
        source = self.root / 'photo.webp'
        exif = Image.Exif()
        exif[315] = 'Private Author'
        Image.new('RGB', (80, 61), '#236778').save(source, format='WEBP', exif=exif)
        output, _ = self.edit(source, {'changes': {'exif:315': 'Novo Autor'}})
        with Image.open(source) as original, Image.open(output) as edited:
            self.assertEqual(original.tobytes(), edited.tobytes())
            self.assertEqual(edited.getexif()[315], 'Novo Autor')
        clean, _ = self.edit(source, {'removeAll': True}, name='clean')
        with Image.open(clean) as image:
            self.assertFalse(image.getexif())

    def test_tiff_structural_fields_read_only(self):
        source = self.root / 'photo.tiff'
        Image.new('RGB', (80, 60), '#467890').save(source, tiffinfo={315: 'Private Author'})
        inspected = tools.inspect(source)
        width = next(r for r in inspected['metadata'] if r['key'] == 'exif:256')
        self.assertFalse(width['editable'])
        self.assertFalse(width['removable'])
        output, _ = self.edit(source, {'removeAll': True})
        with Image.open(source) as original, Image.open(output) as edited:
            self.assertEqual(original.tobytes(), edited.tobytes())
            self.assertNotIn(315, edited.tag_v2)

    def test_tiff_nested_exif_edits_preserve_other_metadata_and_pixels(self):
        source = self.root / 'nested.tiff'
        camera = {36867: '2026:10:22 10:11:12', 42036: 'Keep Lens', 40965: {1: 'R98'}}
        gps = {1: 'N', 2: (1.0, 2.0, 3.0)}
        tags = {315: 'Private Author', 282: 150.0, 283: 150.0, 296: 2,
                34665: camera, 34853: gps}
        Image.linear_gradient('L').convert('RGB').save(source, tiffinfo=tags)
        with Image.open(source) as image:
            pixels = image.tobytes()
            structural = {key: image.tag_v2[key] for key in (256, 257, 258, 262, 277, 282, 283, 284, 296)}
        self.assertTrue(any(item['key'] == 'exifSub:36867' and item['removable']
                            for item in tools.inspect(source)['metadata']))
        for name, remove in [('top-level', []), ('selective', ['exifSub:36867'])]:
            with self.subTest(operation=name):
                output, _ = self.edit(source, {'changes': {'exif:315': 'New Author'}, 'remove': remove}, name=name)
                # Corrupt offsets can look like empty metadata while reporting success.
                with warnings.catch_warnings():
                    warnings.simplefilter('error')
                    with Image.open(output) as image:
                        self.assertEqual({key: image.tag_v2[key] for key in structural}, structural)
                        self.assertEqual(image.getexif()[315], 'New Author')
                        camera_after = image.getexif().get_ifd(34665)
                        if remove:
                            self.assertNotIn(36867, camera_after)
                        else:
                            self.assertEqual(camera_after[36867], camera[36867])
                        self.assertEqual(camera_after[42036], 'Keep Lens')
                        self.assertEqual(image.getexif().get_ifd(40965), {1: 'R98'})
                        self.assertEqual(image.getexif().get_ifd(34853), gps)
                        self.assertEqual(image.tobytes(), pixels)
        cleaned, _ = self.edit(source, {'removeAll': True}, name='cleaned-nested')
        with Image.open(cleaned) as image:
            self.assertEqual(image.tobytes(), pixels)
            self.assertNotIn(34665, image.getexif())
            self.assertNotIn(34853, image.getexif())

    def test_pdf_edit_removes_metadata_preserving_pages(self):
        source = self.root / 'document.pdf'
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=300)
        writer.add_metadata({'/Title': 'Private Title', '/Author': 'Private Author'})
        with source.open('wb') as stream:
            writer.write(stream)
        output, _ = self.edit(source, {'changes': {'pdf:Title': 'Novo Título'}, 'remove': ['pdf:Author']})
        reader = PdfReader(output)
        self.assertEqual(len(reader.pages), 1)
        self.assertEqual(reader.metadata.title, 'Novo Título')
        self.assertNotIn('/Author', reader.metadata)
        clean, _ = self.edit(source, {'removeAll': True}, name='clean')
        self.assertFalse(PdfReader(clean).metadata)

    def test_encrypted_pdf_rejected(self):
        source = self.root / 'locked.pdf'
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=300)
        writer.encrypt('private')
        writer.write(source)
        with self.assertRaises(tools.ToolError):
            tools.inspect(source)
        self.assertFalse((self.root / 'result.pdf').exists())

    def test_ooxml_edit_remove_all_does_not_change_document(self):
        from docx import Document
        source = self.root / 'document.docx'
        document = Document()
        document.add_paragraph('The document content stays intact.')
        document.core_properties.title = 'Private Title'
        document.core_properties.author = 'Private Author'
        document.save(source)
        output, _ = self.edit(source, {'changes': {'core:title': 'Novo Título'}, 'remove': ['core:creator']})
        self.assertEqual(Document(output).core_properties.title, 'Novo Título')
        self.assertEqual(Document(output).core_properties.author, '')
        clean, _ = self.edit(source, {'removeAll': True}, name='clean')
        self.assertEqual(Document(clean).core_properties.title, '')
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(clean) as after:
            self.assertEqual(before.read('word/document.xml'), after.read('word/document.xml'))

    def test_reject_edit_structural_fields_and_preserve_existing_destination(self):
        source = self.jpg()
        target = self.root / 'result.jpg'
        with self.assertRaises(tools.ToolError):
            tools.execute({'action': 'metadata', 'inputs': [str(source)], 'output': str(target), 'options': {'changes': {'image:width': '1'}}})
        self.assertFalse(target.exists())
        target.write_bytes(b'Existing file')
        with self.assertRaises(tools.ToolError):
            tools.execute({'action': 'compress', 'inputs': [str(source)], 'output': str(target)})
        self.assertEqual(target.read_bytes(), b'Existing file')

    def test_zip_names_and_lossless_contents(self):
        directory = self.root / 'other'
        directory.mkdir()
        first, second = self.root / 'sample.txt', directory / 'sample.txt'
        first.write_text('a' * 20000)
        second.write_text('b' * 20000)
        output = self.root / 'archive.zip'
        result = tools.execute({'action': 'compress', 'inputs': [str(first), str(second)], 'output': str(output)})
        self.assertGreater(result['details']['savedPercent'], 0)
        with zipfile.ZipFile(output) as archive:
            self.assertEqual(archive.namelist(), ['sample.txt', 'sample (2).txt'])
            self.assertEqual(archive.read('sample.txt'), first.read_bytes())
            self.assertEqual(archive.read('sample (2).txt'), second.read_bytes())

    def test_zip_presets_select_real_streaming_compression_levels(self):
        source = self.root / 'payload.txt'
        content = (b'compressible payload ' * 100 + bytes(range(256))) * 100
        source.write_bytes(content)
        for preset, level in [('balanced', 6), ('strong', 9)]:
            with self.subTest(preset=preset):
                output = self.root / (preset + '.zip')
                # Wrap the real compressor: verify the level actually used by
                # the streaming writer, then decompress and compare every byte.
                with mock.patch.object(zlib, 'compressobj', wraps=zlib.compressobj) as compressor:
                    tools.execute({'action': 'compress', 'inputs': [str(source)], 'output': str(output),
                                   'options': {'mode': 'zip', 'preset': preset}})
                    self.assertEqual(compressor.call_count, 1)
                    self.assertEqual(compressor.call_args.args[0], level)
                with zipfile.ZipFile(output) as archive:
                    self.assertEqual(archive.read(source.name), content)
                    self.assertLess(archive.getinfo(source.name).compress_size, len(content))
                self.assertEqual(source.read_bytes(), content)

    def test_cli_wire_is_ascii_and_preserves_unicode_metadata(self):
        source = self.root / 'unicode.png'
        info = PngInfo()
        author = '照片 🍈'
        info.add_itxt('Author', author)
        Image.new('RGB', (20, 20), '#234567').save(source, pnginfo=info)
        request = self.root / 'request.json'
        request.write_text(json.dumps({'action': 'inspect', 'input': str(source)}, ensure_ascii=False), encoding='utf-8')
        program = ("import runpy,sys; sys.stdout.reconfigure(encoding='ascii',errors='strict'); "
                   "sys.argv=[" + repr(str(HELPER)) + ", '--request', " + repr(str(request)) + "]; "
                   "runpy.run_path(" + repr(str(HELPER)) + ",run_name='__main__')")
        child = subprocess.run([sys.executable, '-I', '-c', program], capture_output=True, check=True, timeout=10)
        wire = child.stdout.decode('ascii')
        result = json.loads(wire)
        self.assertTrue(result['ok'])
        value = next(item['value'] for item in result['metadata'] if item['key'] == 'text:Author')
        self.assertEqual(value, author)
        self.assertIn('\\u', wire)
        # Failure responses also contain localized text and must use ASCII JSON.
        request.write_text('{}', encoding='utf-8')
        child = subprocess.run([sys.executable, '-I', '-c', program], capture_output=True, timeout=10)
        self.assertEqual(child.returncode, 1)
        failure = json.loads(child.stdout.decode('ascii'))
        self.assertFalse(failure['ok'])
        self.assertTrue(failure['error'])

    def test_archive_names_are_portable_basenames(self):
        ordinary = ['photo.jpg', 'Meu arquivo.pdf', 'Maracujá.txt', 'com10.txt', 'company.docx']
        for name in ordinary:
            self.assertEqual(tools.archive_name(name), name)
        cases = {
            '..\\escape.txt': '_escape.txt',
            'C:\\private\\file.txt': 'C__private_file.txt',
            '../escape.txt': '_escape.txt',
            '...': 'arquivo',
            ' . . ': 'arquivo',
            'report.txt. ': 'report.txt',
            'CON': '_CON',
            'nul.txt': '_nul.txt',
            'PrN .txt': '_PrN .txt',
            'COM1.pdf': '_COM1.pdf',
            'lpt9': '_lpt9',
            'AUX.data': '_AUX.data',
            'bad:<name>?.txt': 'bad__name__.txt',
            'control\x01\x7f.txt': 'control__.txt',
        }
        for original, expected in cases.items():
            with self.subTest(name=original):
                safe = tools.archive_name(original)
                self.assertEqual(safe, expected)
                self.assertNotIn('/', safe)
                self.assertNotIn('\\', safe)
                self.assertNotIn(':', safe)
                self.assertFalse(safe.startswith('.'))
                self.assertFalse(safe.endswith(('.', ' ')))
        # Simulate POSIX-only names without creating invalid Windows files.
        content = self.root / 'payload.txt'
        content.write_bytes(b'fixture bytes')
        class NamedSource:
            def __init__(self, name):
                self.name = name
            def open(self, mode):
                return content.open(mode)
        archive = self.root / 'portable.zip'
        note = tools.compress_zip([NamedSource('CON.txt'), NamedSource('_con.txt'), NamedSource('..\\escape.txt')], archive, {})
        with zipfile.ZipFile(archive) as result:
            self.assertEqual(result.namelist(), ['_CON.txt', '_con (2).txt', '_escape.txt'])
            self.assertTrue(all(result.read(name) == b'fixture bytes' for name in result.namelist()))
        self.assertIn('nomes foram ajustados', note)

    def test_image_compression_target_truthful_and_no_growth(self):
        source = self.root / 'unoptimized.jpg'
        image = Image.effect_noise((512, 512), 25).convert('RGB')
        image.save(source, quality=100)
        output, result = self.edit(source, {'preset': 'strong', 'targetPercent': 20, 'maxEdge': 256}, action='compress')
        self.assertLessEqual(output.stat().st_size, source.stat().st_size)
        self.assertTrue(result['details']['targetMet'])
        with Image.open(output) as image:
            self.assertLessEqual(max(image.size), 256)
        self.assertEqual(tools.compression_target({'targetPercent': 30}, 1000), 700)

    def test_impossible_target_reported_and_optimized_original_copied(self):
        source = self.root / 'already.png'
        Image.new('RGB', (10, 10), 'white').save(source, optimize=True)
        output, result = self.edit(source, {'targetBytes': 1}, action='compress')
        self.assertFalse(result['details']['targetMet'])
        self.assertLessEqual(output.stat().st_size, source.stat().st_size)

    def test_resize_is_honored_even_when_optimized_source_was_smaller(self):
        from PIL import ImageDraw
        source = self.root / 'tiny-optimized.png'
        image = Image.new('RGB', (100, 100), 'black')
        ImageDraw.Draw(image).line((0, 0, 99, 99), fill='white', width=1)
        image.save(source, optimize=True)
        output, result = self.edit(source, {'maxEdge': 99}, action='compress')
        with Image.open(output) as resized:
            self.assertEqual(resized.size, (99, 99))
        self.assertGreater(output.stat().st_size, source.stat().st_size)
        self.assertLess(result['details']['savedPercent'], 0)
        self.assertEqual(result['details']['outputBytes'], output.stat().st_size)
        self.assertIn('Não houve redução', result['details']['note'])

    def test_pdf_image_dimensions_checked_before_decode(self):
        # The XObject is only a dictionary, so no image samples are allocated.
        page = {'/Resources': {'/XObject': {'/Im1': {'/Subtype': '/Image', '/Width': 100000, '/Height': 100000}}}}
        with self.assertRaises(tools.ToolError):
            tools.check_pdf_image_bounds(page)

    def test_pdf_resize_is_retained(self):
        source = self.root / 'scan.pdf'
        Image.new('RGB', (200, 200), '#234567').save(source, format='PDF')
        output, _ = self.edit(source, {'maxEdge': 99}, action='compress')
        reader = PdfReader(output)
        self.assertEqual(len(reader.pages), 1)
        self.assertLessEqual(max(reader.pages[0].images[0].image.size), 99)

    @unittest.skipUnless(os.name == 'posix', 'POSIX worker memory budget')
    def test_memory_budget_applies_only_when_worker_explicitly_starts(self):
        # The disposable child checks real limits. Importing the library must
        # never limit the test runner or GUI; only its CLI worker opts in.
        program = ("import importlib.util,json,resource; "
                   "before=resource.getrlimit(resource.RLIMIT_AS); "
                   "s=importlib.util.spec_from_file_location('file_tools_budget'," + repr(str(HELPER)) + "); "
                   "m=importlib.util.module_from_spec(s); s.loader.exec_module(m); "
                   "imported=resource.getrlimit(resource.RLIMIT_AS); "
                   "m.limit_worker_memory(); after=resource.getrlimit(resource.RLIMIT_AS); "
                   "resource.setrlimit(resource.RLIMIT_AS,(1024*1024*1024,after[1])); "
                   "m.limit_worker_memory(); stricter=resource.getrlimit(resource.RLIMIT_AS); "
                   "print(json.dumps([before,imported,after,stricter]))")
        child = subprocess.run([sys.executable, '-I', '-c', program], capture_output=True, check=True, timeout=10)
        before, imported, after, stricter = json.loads(child.stdout)
        self.assertEqual(before, imported)
        self.assertLessEqual(after[0], 2 * 1024 * 1024 * 1024)
        self.assertEqual(after[1], before[1])
        self.assertEqual(stricter[0], 1024 * 1024 * 1024)

    def test_removal_wins_over_stale_edit_drafts(self):
        source = self.jpg()
        clean, _ = self.edit(source, {'removeAll': True, 'changes': {'exif:315': 'Stale draft'}})
        with Image.open(clean) as image:
            self.assertNotIn(315, image.getexif())
        selected, _ = self.edit(source, {'remove': ['exif:315'], 'changes': {'exif:315': 'Stale draft'}}, name='selected')
        with Image.open(selected) as image:
            self.assertNotIn(315, image.getexif())

    def test_non_file_input_and_malformed_xml_rejected(self):
        with self.assertRaises(tools.ToolError):
            tools.source_path(str(self.root))
        with self.assertRaises(tools.ToolError):
            tools.safe_xml(b'<!DOCTYPE doc [<!ENTITY key "value">]><doc/>')


if __name__ == '__main__':
    unittest.main(verbosity=2)
