#!/usr/bin/python3
"""Sequential end-to-end checks using only generated, disposable local files.

Usage: python3 -I tests/expanded_formats_check.py build/mandarina
No personal documents, network resources, desktop session or existing app process
are used. Each subprocess is awaited before the next conversion starts.
"""

import array
import hashlib
import html
import importlib.util
from html.parser import HTMLParser
import io
import json
import math
import pathlib
import re
import struct
import subprocess
import sys
import tempfile
import unicodedata
import wave
import xml.etree.ElementTree as ET
import zipfile
from xml.sax.saxutils import escape as xml_escape

from docx import Document
from openpyxl import load_workbook
import pdfplumber
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle


_runtime_spec = importlib.util.spec_from_file_location("maracuja_test_runtime", pathlib.Path(__file__).with_name("test_runtime.py"))
_runtime_module = importlib.util.module_from_spec(_runtime_spec)
_runtime_spec.loader.exec_module(_runtime_module)
APP = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else
                   _runtime_module.default_application(pathlib.Path(__file__).resolve().parents[1])).resolve()
RUNTIME = _runtime_module.TestRuntime(APP)
DOCUMENT_FORMATS = {"pdf", "docx", "odt", "rtf", "txt", "md", "html", "epub", "png", "jpg"}
PDF_FORMATS = {"png", "jpg", "txt", "pdf", "docx", "md", "html", "epub", "pptx", "xlsx"}
AUDIO_CODECS = {
    "mp3": "mp3", "m4a": "aac", "wav": "pcm_s24le", "flac": "flac",
    "ogg": "vorbis", "opus": "opus", "aiff": "pcm_s24be", "wma": "wmav2",
}
MARKER = "TESTE MARACUJA NUMERO 5729"
UNICODE_WORDS = ("Maracujá", "ação", "café", "São Paulo")
UNICODE_SENTENCE = "Maracujá, ação e café — São Paulo."
LAST_MARKER = "FIM DO DOCUMENTO 8614"
checks = []
current_check = "initialization"


def require(condition, reason):
    if not condition:
        raise AssertionError(reason)


def start_check(name):
    global current_check
    current_check = name
    print("Checking " + name, file=sys.stderr, flush=True)


def passed():
    checks.append(current_check)


def sha256(path):
    digest = hashlib.sha256()
    with pathlib.Path(path).open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def command(arguments, timeout=120):
    # Generated fixtures only. Do not echo arguments or captured document text.
    result = subprocess.run(list(map(str, arguments)), stdin=subprocess.DEVNULL,
                            stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                            timeout=timeout, check=False, **RUNTIME.process_options())
    return result


def convert(fmt, sources, destination=None, success=True):
    arguments = [APP, "--convert", fmt]
    if destination is not None:
        arguments.extend(["--output", destination])
    result = command([*arguments, "--", *sources])
    require(result.returncode == (0 if success else 1), "unexpected CLI exit status")
    try:
        report = json.loads(result.stdout)
    except (ValueError, UnicodeError) as error:
        raise AssertionError("conversion did not return JSON") from error
    require(isinstance(report.get("outputs"), list) and isinstance(report.get("errors"), list),
            "invalid conversion report")
    require(not report.get("cancelled"), "conversion unexpectedly cancelled")
    if success:
        require(not report["errors"] and bool(report["outputs"]), "conversion reported errors")
    else:
        require(bool(report["errors"]) and not report["outputs"], "failed conversion published output")
    paths = list(map(pathlib.Path, report["outputs"]))
    destinations = ({destination.resolve()} if destination is not None else
                    {pathlib.Path(source).resolve().parent for source in sources})
    for path in paths:
        require(path.is_file() and not path.is_symlink(), "missing, directory or symbolic output")
        require(path.resolve().parent in destinations, "output escaped requested directory")
        require(path.stat().st_size > 0, "empty output file")
    return paths, report


def convert_one(fmt, source, destination):
    paths, _ = convert(fmt, [source], destination)
    require(len(paths) == 1, "single conversion produced unexpected output count")
    return paths[0]


def assert_page_images(paths, source, destination, fmt, count, original_names=True):
    require(len(paths) == count and len(set(paths)) == count, "document rendering lost or duplicated pages")
    for index, path in enumerate(paths, 1):
        require(path.is_file() and path.resolve().parent == destination.resolve(),
                "rendered page was not saved directly in destination")
        if original_names:
            suffix = "" if count == 1 else "-pagina-" + str(index).zfill(len(str(count)))
            require(path.name == source.stem + "-convertido" + suffix + "." + fmt,
                    "rendered page name or page order is incorrect")
        with Image.open(path) as image:
            require(image.format == ("PNG" if fmt == "png" else "JPEG"),
                    "rendered page uses the wrong image format")
            require(300 < max(image.size) <= 4096, "rendered page dimensions invalid")
            require(image.convert("L").getextrema()[0] < 200, "rendered page is blank")


def normalized(text):
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text)).strip()


def assert_text(text, include_end=False):
    text = normalized(text)
    require(MARKER in text, "document content marker was lost")
    require(all(word in text for word in UNICODE_WORDS), "Unicode document text was lost")
    if include_end:
        require(LAST_MARKER in text, "last page content was lost")


class HtmlText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.hidden_depth = 0
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden_depth += 1

    def handle_endtag(self, tag):
        if tag in {"script", "style"} and self.hidden_depth:
            self.hidden_depth -= 1
        if tag in {"p", "div", "li", "h1", "h2", "h3", "tr", "br"}:
            self.parts.append(" ")

    def handle_data(self, data):
        if not self.hidden_depth:
            self.parts.append(data)


def html_text(text):
    parser = HtmlText()
    parser.feed(text)
    return " ".join(parser.parts)


def rtf_text(data):
    """Read plain text independently from common RTF control words.

    This only serves generated text fixtures; it is not a production importer.
    It handles Unicode escapes/fallbacks, CP1252 escapes and ignored destination
    groups, so a file merely renamed .rtf cannot pass as a valid conversion.
    """
    require(data.startswith(b"{\\rtf"), "RTF signature missing")
    text = data.decode("latin-1")
    ignored = {"fonttbl", "colortbl", "stylesheet", "info", "pict", "object", "header",
               "footer", "filetbl", "generator", "listtable", "listoverridetable",
               "datastore", "themedata", "colorschememapping", "latentstyles", "xmlnstbl"}
    tokens = re.compile(r"\\([a-zA-Z]+)(-?\d+)? ?|\\'([0-9a-fA-F]{2})|\\([^a-zA-Z])|([{}])|([^\\{}]+)")
    stack = []
    uc, ignorable, skip = 1, False, 0
    parts = []
    for token in tokens.finditer(text):
        word, number, hex_byte, symbol, brace, literal = token.groups()
        if brace == "{":
            stack.append((uc, ignorable))
        elif brace == "}":
            if stack:
                uc, ignorable = stack.pop()
        elif word:
            if word in ignored:
                ignorable = True
            elif word == "uc":
                uc = max(0, int(number or "1"))
            elif not ignorable:
                if word == "u":
                    parts.append(chr(int(number or "0") & 0xFFFF))
                    skip = uc
                elif word in {"par", "line", "tab"}:
                    parts.append(" ")
                elif word in {"emdash", "endash", "bullet"}:
                    parts.append({"emdash": "—", "endash": "–", "bullet": "•"}[word])
        elif symbol:
            if symbol == "*":
                ignorable = True
            elif not ignorable and symbol in {"{", "}", "\\", "~", "_"}:
                if skip:
                    skip -= 1
                else:
                    parts.append({"~": " ", "_": "-"}.get(symbol, symbol))
        elif hex_byte and not ignorable:
            if skip:
                skip -= 1
            else:
                parts.append(bytes.fromhex(hex_byte).decode("cp1252", errors="replace"))
        elif literal and not ignorable:
            # Formatting newlines in the RTF source are not document characters.
            literal = literal.replace("\r", "").replace("\n", "")
            skipped = min(skip, len(literal))
            skip -= skipped
            parts.append(literal[skipped:])
    return "".join(parts)


def read_document(path, fmt):
    if fmt in {"txt", "md"}:
        return path.read_text(encoding="utf-8-sig")
    if fmt == "html":
        return html_text(path.read_text(encoding="utf-8-sig"))
    if fmt == "pdf":
        require(path.read_bytes()[:5] == b"%PDF-", "PDF signature missing")
        with pdfplumber.open(path) as document:
            require(bool(document.pages), "PDF has no pages")
            pages = []
            for index, page in enumerate(document.pages, 1):
                # Cambria uses /ActualText for composed Unicode characters.
                # pdfplumber ignores those spans; use Poppler for those pages,
                # retaining ordinary layout extraction on the remaining ones.
                if any(b"/ActualText" in stream.get_data() for stream in page.page_obj.contents):
                    result = command([RUNTIME.tool("pdftotext"), "-raw", "-enc", "UTF-8", "-nopgbrk",
                                      "-f", str(index), "-l", str(index), path, "-"])
                    require(result.returncode == 0, "PDF text extraction failed")
                    pages.append(result.stdout.decode("utf-8"))
                else:
                    pages.append(page.extract_text() or "")
            return "\n".join(pages)
    if fmt == "docx":
        document = Document(path)
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if fmt == "odt":
        with zipfile.ZipFile(path) as archive:
            require(archive.read("mimetype") == b"application/vnd.oasis.opendocument.text",
                    "ODT mimetype missing")
            body = ET.fromstring(archive.read("content.xml"))
            return " ".join(body.itertext())
    if fmt == "rtf":
        return rtf_text(path.read_bytes())
    if fmt == "epub":
        with zipfile.ZipFile(path) as archive:
            require(archive.read("mimetype") == b"application/epub+zip", "EPUB mimetype missing")
            require("META-INF/container.xml" in archive.namelist(), "EPUB container missing")
            container = ET.fromstring(archive.read("META-INF/container.xml"))
            rootfiles = [node.attrib.get("full-path") for node in container.iter()
                         if node.tag.rsplit("}", 1)[-1] == "rootfile"]
            require(bool(rootfiles) and all(name in archive.namelist() for name in rootfiles),
                    "EPUB package missing")
            texts = [html_text(archive.read(name).decode("utf-8")) for name in archive.namelist()
                     if pathlib.PurePosixPath(name).suffix.lower() in {".xhtml", ".html", ".htm"}]
            require(bool(texts), "EPUB has no readable document")
            return "\n".join(texts)
    raise AssertionError("unknown document validator")


def make_odt(path, paragraphs):
    content = ("<?xml version='1.0' encoding='UTF-8'?>"
               "<office:document-content xmlns:office='urn:oasis:names:tc:opendocument:xmlns:office:1.0' "
               "xmlns:text='urn:oasis:names:tc:opendocument:xmlns:text:1.0' office:version='1.2'>"
               "<office:body><office:text>" +
               "".join("<text:p>" + xml_escape(paragraph) + "</text:p>" for paragraph in paragraphs) +
               "</office:text></office:body></office:document-content>")
    manifest = ("<?xml version='1.0' encoding='UTF-8'?>"
                "<manifest:manifest xmlns:manifest='urn:oasis:names:tc:opendocument:xmlns:manifest:1.0' "
                "manifest:version='1.2'><manifest:file-entry manifest:full-path='/' "
                "manifest:media-type='application/vnd.oasis.opendocument.text'/>"
                "<manifest:file-entry manifest:full-path='content.xml' manifest:media-type='text/xml'/>"
                "</manifest:manifest>")
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", b"application/vnd.oasis.opendocument.text",
                         compress_type=zipfile.ZIP_STORED)
        archive.writestr("content.xml", content.encode("utf-8"), compress_type=zipfile.ZIP_DEFLATED)
        archive.writestr("META-INF/manifest.xml", manifest.encode("utf-8"),
                         compress_type=zipfile.ZIP_DEFLATED)


def make_rtf(path, paragraphs):
    def escaped(text):
        return "".join("\\" + char if char in "{}\\" else char if ord(char) < 128 else
                       "\\u" + str(ord(char) if ord(char) < 32768 else ord(char) - 65536) + "?"
                       for char in text)
    path.write_bytes(("{\\rtf1\\ansi\\ansicpg1252\\uc1 " +
                      "\\par\n".join(map(escaped, paragraphs)) + "}").encode("ascii"))


def make_text_pdf(path):
    document = canvas.Canvas(str(path), pagesize=A4)
    for page in range(2):
        document.setFillColor(colors.HexColor("#6842a6" if page == 0 else "#b68a00"))
        document.rect(48, 720, 480, 50, fill=1, stroke=0)
        document.setFillColor(colors.black)
        document.setFont("Helvetica", 13)
        document.drawString(50, 680, MARKER)
        document.drawString(50, 650, UNICODE_SENTENCE)
        document.drawString(50, 620, "Pagina " + str(page + 1))
        if page == 1:
            document.drawString(50, 590, LAST_MARKER)
        document.showPage()
    document.save()


def make_table_pdf(path):
    table = Table([["Produto", "Quantidade", "Valor"],
                   ["Maracuja", "3", "12,50"], ["Cafe", "2", "8,00"]],
                  colWidths=[190, 100, 100], rowHeights=[32, 32, 32])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ece4fa")),
        ("FONTSIZE", (0, 0), (-1, -1), 12),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    SimpleDocTemplate(str(path), pagesize=A4).build([table])


def make_scanned_pdf(path):
    image = Image.new("RGB", (900, 500), "white")
    # A genuinely empty scan remains an error even after OCR is available.
    document = canvas.Canvas(str(path), pagesize=A4)
    document.drawImage(ImageReader(image), 50, 400, width=490, height=272)
    document.save()
    image.close()


def make_audio(path):
    rate = 16000
    samples = [int(7000 * math.sin(2 * math.pi * 880 * frame / rate))
               for frame in range(int(rate * 0.7))]
    with wave.open(str(path), "wb") as destination:
        destination.setnchannels(1)
        destination.setsampwidth(2)
        destination.setframerate(rate)
        destination.writeframes(struct.pack("<" + "h" * len(samples), *samples))


def assert_audio(path, codec):
    result = command([RUNTIME.tool("ffprobe"), "-v", "error", "-select_streams", "a:0",
                      "-show_entries", "stream=codec_name,channels,sample_rate:format=duration",
                      "-of", "json", "-i", path], timeout=20)
    require(result.returncode == 0, "output is not readable audio")
    probe = json.loads(result.stdout)
    streams = probe.get("streams", [])
    require(len(streams) == 1 and streams[0].get("codec_name") == codec,
            "audio encoder/container does not match requested format")
    require(streams[0].get("channels") == 1, "mono audio was unexpectedly changed to stereo")
    require(0.5 <= float(probe.get("format", {}).get("duration", 0)) <= 1.4,
            "audio duration was lost or truncated")
    decoded = command([RUNTIME.tool("ffmpeg"), "-nostdin", "-v", "error", "-i", path,
                       "-map", "0:a:0", "-ac", "1", "-ar", "16000", "-f", "s16le",
                       "-c:a", "pcm_s16le", "pipe:1"], timeout=20)
    require(decoded.returncode == 0 and len(decoded.stdout) > 16000,
            "audio output contains no decodable samples")
    samples = array.array("h")
    samples.frombytes(decoded.stdout)
    if sys.byteorder != "little":
        samples.byteswap()
    # Inspect the middle of the generated tone, away from codec priming/padding.
    middle = samples[1600:8000]
    require(len(middle) == 6400, "audio output is too short")
    rms = math.sqrt(sum(sample * sample for sample in middle) / len(middle))
    crossings = sum((before < 0) != (after < 0) for before, after in zip(middle, middle[1:]))
    frequency = crossings * 16000 / (2 * len(middle))
    require(1000 < rms < 15000 and 800 < frequency < 960,
            "audio samples do not preserve the generated tone")


def main():
    start_check("advertised_document_pdf_and_audio_capabilities")
    require(APP.is_file(), "converter binary missing")
    result = command([APP, "--list-formats"], timeout=30)
    require(result.returncode == 0, "capability listing failed")
    capabilities = json.loads(result.stdout)
    require(DOCUMENT_FORMATS <= set(capabilities.get("documents", [])),
            "required document formats unavailable")
    require(PDF_FORMATS <= set(capabilities.get("pdf", [])), "required PDF formats unavailable")
    require(set(AUDIO_CODECS) <= set(capabilities.get("audio", [])), "required audio encoders unavailable")
    passed()

    with tempfile.TemporaryDirectory(prefix="maracuja-expanded-check-") as temporary:
        root = pathlib.Path(temporary)
        inputs = root / "entradas privadas"
        outputs = root / "saidas"
        inputs.mkdir(mode=0o700)
        outputs.mkdir(mode=0o700)
        hashes = {}

        def track(path):
            hashes[path] = sha256(path)
            return path

        paragraphs = [MARKER, UNICODE_SENTENCE, "Documento local gerado para validar conversoes."]
        text_source = inputs / "texto ação com espaços.txt"
        lines = paragraphs + ["Linha %03d - texto de teste com acentos: ação e café." % number
                              for number in range(1, 171)] + [LAST_MARKER]
        text_source.write_text("\n".join(lines) + "\n", encoding="utf-8")
        track(text_source)

        start_check("text_mentioning_pdf_header_remains_document")
        header_note = inputs / "cabecalho.txt"
        header_note.write_text("Cabeçalho PDF: %PDF-1.7\n" + "\n".join(paragraphs), encoding="utf-8")
        track(header_note)
        assert_text(read_document(convert_one("docx", header_note, outputs), "docx"))
        passed()

        document_results = {}
        for fmt in ["pdf", "docx", "odt", "rtf", "txt", "md", "html", "epub"]:
            start_check("unicode_txt_to_" + fmt)
            path = convert_one(fmt, text_source, outputs)
            assert_text(read_document(path, fmt), include_end=True)
            document_results[fmt] = path
            passed()

        with pdfplumber.open(document_results["pdf"]) as document:
            text_pages = len(document.pages)
        require(text_pages >= 2, "multipage fixture did not paginate")
        for fmt in ["png", "jpg"]:
            start_check("txt_all_pages_to_" + fmt)
            pages, _ = convert(fmt, [text_source], outputs)
            assert_page_images(pages, text_source, outputs, fmt, text_pages)
            passed()
            start_check("document_page_" + fmt + "_collisions_preserve_every_existing_file")
            previous = {page: sha256(page) for page in pages}
            repeated, _ = convert(fmt, [text_source], outputs)
            assert_page_images(repeated, text_source, outputs, fmt, text_pages, original_names=False)
            require(set(repeated).isdisjoint(pages), "page collision reused an existing image")
            require(all(sha256(page) == digest for page, digest in previous.items()),
                    "page collision changed an existing image")
            passed()
            start_check("single_document_page_" + fmt + "_without_output_beside_original")
            single, _ = convert(fmt, [header_note])
            assert_page_images(single, header_note, inputs, fmt, 1)
            passed()

        docx_source = inputs / "fonte.docx"
        document = Document()
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        document.save(docx_source)
        odt_source = inputs / "fonte.odt"
        make_odt(odt_source, paragraphs)
        rtf_source = inputs / "fonte.rtf"
        make_rtf(rtf_source, paragraphs)
        md_source = inputs / "fonte.md"
        md_source.write_text("# " + MARKER + "\n\n" + UNICODE_SENTENCE + "\n", encoding="utf-8")
        html_source = inputs / "fonte.html"
        html_source.write_text("<!doctype html><html><head><meta charset='utf-8'></head><body><h1>" +
                               html.escape(MARKER) + "</h1><p>" + html.escape(UNICODE_SENTENCE) +
                               "</p></body></html>", encoding="utf-8")
        for source in [docx_source, odt_source, rtf_source, md_source, html_source]:
            track(source)
            for fmt in ["txt", "pdf"]:
                start_check(source.suffix[1:] + "_source_to_" + fmt + "_retains_text")
                path = convert_one(fmt, source, outputs)
                assert_text(read_document(path, fmt))
                passed()
                if source == docx_source and fmt == "pdf":
                    start_check("docx_pdf_actual_text_to_docx_retains_text")
                    assert_text(read_document(convert_one("docx", path, outputs), "docx"))
                    passed()

        text_pdf = inputs / "pdf com duas páginas.pdf"
        make_text_pdf(text_pdf)
        track(text_pdf)
        start_check("pdf_detected_by_content_supports_document_export")
        renamed_pdf = inputs / "documento.bin"
        renamed_pdf.write_bytes(text_pdf.read_bytes())
        track(renamed_pdf)
        assert_text(read_document(convert_one("docx", renamed_pdf, outputs), "docx"), include_end=True)
        passed()
        for fmt in ["docx", "md", "html", "epub"]:
            start_check("pdf_text_to_" + fmt)
            path = convert_one(fmt, text_pdf, outputs)
            assert_text(read_document(path, fmt), include_end=True)
            passed()

        start_check("pdf_table_to_real_spreadsheet_cells")
        table_pdf = inputs / "tabela.pdf"
        make_table_pdf(table_pdf)
        track(table_pdf)
        spreadsheet = convert_one("xlsx", table_pdf, outputs)
        workbook = load_workbook(spreadsheet, read_only=True, data_only=True)
        rows = [tuple(str(value).strip() if value is not None else "" for value in row)
                for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        require(any("Produto" in row and "Quantidade" in row for row in rows),
                "spreadsheet headers missing")
        require(any("Maracuja" in row and "3" in row and
                    ("12,50" in row or "12.5" in row) for row in rows),
                "spreadsheet table cells were flattened or lost")
        require(any("Cafe" in row and "2" in row for row in rows), "second table row missing")
        passed()

        start_check("pdf_without_table_rejects_spreadsheet_export")
        convert("xlsx", [text_pdf], outputs, success=False)
        passed()

        start_check("pdf_to_pptx_has_image_for_every_page")
        slides_path = convert_one("pptx", text_pdf, outputs)
        slides = Presentation(slides_path)
        require(len(slides.slides) == 2, "presentation page count differs from PDF")
        slide_hashes = []
        for slide in slides.slides:
            pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
            require(len(pictures) >= 1, "PDF page missing from slide")
            image_bytes = pictures[0].image.blob
            with Image.open(io.BytesIO(image_bytes)) as image:
                require(max(image.size) >= 500, "slide page image resolution too small")
            slide_hashes.append(hashlib.sha256(image_bytes).hexdigest())
        require(len(set(slide_hashes)) == 2, "different PDF pages became duplicate slides")
        passed()

        scanned_pdf = inputs / "digitalizado sem texto.pdf"
        make_scanned_pdf(scanned_pdf)
        track(scanned_pdf)
        for fmt in ["txt", "docx", "md", "html", "epub"]:
            start_check("blank_scanned_pdf_does_not_publish_empty_" + fmt)
            _, report = convert(fmt, [scanned_pdf], outputs, success=False)
            error_text = " ".join(report["errors"]).lower()
            require("texto" in error_text, "missing explanation for absent recognized text")
            require(not re.search(r"ocr.*(?:ainda|indispon|não.{0,15}dispon)", error_text),
                    "OCR was unavailable instead of examining the empty scan")
            passed()

        audio_source = inputs / "-som ação mono.wav"
        make_audio(audio_source)
        track(audio_source)
        audio_results = {}
        for fmt, codec in AUDIO_CODECS.items():
            start_check("mono_wav_to_" + fmt + "_retains_audio")
            path = convert_one(fmt, audio_source, outputs)
            assert_audio(path, codec)
            audio_results[fmt] = path
            passed()
        for input_format, output_format in [("mp3", "flac"), ("m4a", "wav")]:
            start_check(input_format + "_roundtrip_to_" + output_format)
            path = convert_one(output_format, audio_results[input_format], outputs)
            assert_audio(path, AUDIO_CODECS[output_format])
            passed()

        invalid_audio = inputs / "playlist disfarçada.wav"
        invalid_audio.write_text("#EXTM3U\n" + audio_source.as_uri() + "\n", encoding="utf-8")
        track(invalid_audio)
        start_check("renamed_playlist_is_not_treated_as_audio")
        convert("mp3", [invalid_audio], outputs, success=False)
        passed()

        invalid_document = inputs / "documento quebrado.docx"
        invalid_document.write_bytes(b"This is not an OOXML document.\x00\x00")
        track(invalid_document)
        start_check("invalid_document_does_not_publish_output")
        convert("pdf", [invalid_document], outputs, success=False)
        passed()

        for source, fmt, existing in [(text_source, "docx", document_results["docx"]),
                                      (audio_source, "wav", audio_results["wav"])]:
            start_check(fmt + "_collision_never_overwrites_existing_output")
            previous = sha256(existing)
            new_output = convert_one(fmt, source, outputs)
            require(new_output != existing, "collision reused an existing file")
            require(sha256(existing) == previous, "existing output changed during collision")
            if fmt == "docx":
                assert_text(read_document(new_output, fmt), include_end=True)
            else:
                assert_audio(new_output, AUDIO_CODECS[fmt])
            passed()

        start_check("every_original_file_is_unchanged")
        require(all(path.is_file() and sha256(path) == digest for path, digest in hashes.items()),
                "an original input file was changed or removed")
        passed()
        start_check("successful_and_failed_jobs_clean_temporary_files")
        require(not list(root.rglob(".mandarina-*")) and not list(root.rglob(".maracuja-*")),
                "conversion left temporary files behind")
        require({path.resolve() for path in root.rglob("*") if path.is_dir()} ==
                {inputs.resolve(), outputs.resolve()}, "conversion left an output subdirectory behind")
        passed()

    print(json.dumps({"passed": len(checks), "checks": checks}, indent=2))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        # Keep automatic validation output free of file names and fixture text.
        reason = str(error) if isinstance(error, AssertionError) else type(error).__name__
        print(json.dumps({"failed": current_check, "reason": reason, "passed": len(checks)}))
        sys.exit(1)
