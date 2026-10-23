#!/usr/bin/env python3
"""Cross-platform smoke tests against the complete staged distribution.

Sequential, generated fixtures only. This verifies that packaged dependencies
actually load; GUI/global-drag behavior still requires an interactive desktop.
"""
import hashlib
import json
import math
import os
from pathlib import Path
import struct
import subprocess
import sys
import tempfile
import wave

from PIL import Image, ImageDraw, ImageFont
from docx import Document

app = Path(sys.argv[1]).resolve()
root = app.parent
environment = os.environ.copy()
environment["MARACUJA_RUNTIME_CONFIG"] = str(root / "runtime-paths.json")
environment["PATH"] = str(root / "tools") + os.pathsep + environment.get("PATH", "")
checks = []


def run(arguments, timeout=180):
    return subprocess.run(list(map(str, arguments)), check=True, stdin=subprocess.DEVNULL,
                          stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=environment, timeout=timeout)


def conversion(source, target, output):
    report = json.loads(run([app, "--convert", target, "--output", output, "--", source]).stdout)
    assert not report["errors"] and not report["cancelled"] and len(report["outputs"]) == 1
    path = Path(report["outputs"][0])
    assert path.is_file() and path.stat().st_size > 0 and path.parent.resolve() == output.resolve()
    return path


capabilities = json.loads(run([app, "--list-formats"]).stdout)
assert {"pdf", "docx", "odt", "rtf", "txt", "md", "html", "epub"} <= set(capabilities["documents"])
assert {"txt", "docx"} <= set(capabilities["images_ocr"])
audio = {"mp3": "mp3", "m4a": "aac", "wav": "pcm_s24le", "flac": "flac", "ogg": "vorbis",
         "opus": "opus", "aiff": "pcm_s24be", "wma": "wmav2"}
assert set(audio) <= set(capabilities["audio"])
checks.append("complete_runtime_capabilities")
with tempfile.TemporaryDirectory(prefix="maracuja-distribution-check-") as temporary:
    directory = Path(temporary)
    output = directory / "outputs"
    output.mkdir()
    source = directory / "texto ação.txt"
    source.write_text("MARACUJA 5729\nAção, café e açúcar.\n", encoding="utf-8")
    original = hashlib.sha256(source.read_bytes()).hexdigest()
    for target in ["pdf", "docx", "odt", "rtf", "txt", "md", "html", "epub"]:
        product = conversion(source, target, output)
        if target == "docx":
            assert "MARACUJA 5729" in "\n".join(p.text for p in Document(product).paragraphs)
        checks.append("packaged_document_" + target)
    image_path = directory / "ocr.png"
    fonts = [root / "fonts/DejaVuSans.ttf", Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
             Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/arial.ttf"]
    font = next((ImageFont.truetype(str(path), 80) for path in fonts if path.is_file()), None)
    assert font is not None
    with Image.new("RGB", (1600, 650), "white") as image:
        ImageDraw.Draw(image).text((100, 160), "MARACUJA 5729", fill="black", font=font)
        image.save(image_path)
    for target in ["txt", "docx"]:
        product = conversion(image_path, target, output)
        text = product.read_text(encoding="utf-8") if target == "txt" else "\n".join(p.text for p in Document(product).paragraphs)
        assert "MARACUJA" in text.upper() and "5729" in text
        checks.append("packaged_ocr_" + target)
    tone = directory / "tone.wav"
    samples = [int(6000 * math.sin(2 * math.pi * 880 * i / 16000)) for i in range(11200)]
    with wave.open(str(tone), "wb") as target:
        target.setnchannels(1)
        target.setsampwidth(2)
        target.setframerate(16000)
        target.writeframes(struct.pack("<" + "h" * len(samples), *samples))
    probe = root / "tools" / ("ffprobe.exe" if os.name == "nt" else "ffprobe")
    for target, expected_codec in audio.items():
        product = conversion(tone, target, output)
        report = json.loads(run([probe, "-v", "error", "-select_streams", "a:0", "-show_entries",
                                 "stream=codec_name,channels", "-of", "json", product]).stdout)
        assert report["streams"][0]["codec_name"] == expected_codec
        assert report["streams"][0]["channels"] == 1
        checks.append("packaged_audio_" + target)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == original
    assert not list(directory.rglob(".mandarina-*")) and not list(directory.rglob(".maracuja-*"))
    checks.append("original_and_temporary_file_safety")
print(json.dumps({"passed": len(checks), "checks": checks, "interactive_gesture_tested": False}, indent=2))
